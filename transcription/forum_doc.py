#!/usr/bin/env python3
"""Build the League of Women Voters candidate-forum transcript deliverable (.docx + .pdf).

A recurring deliverable (2022, 2024, 2026 editions): one document per election year,
filed as `<date> Troy School Board Candidate Forum - League of Women Voters.{docx,pdf}`.
The layout reproduces the 2024 edition exactly (title block 20/13/11/9 pt, "Questions
Asked of the Candidates" as a numbered list with [m:ss] timestamps and 9 pt indented
notes, then "Full Transcript" as `[m:ss] Name: text` lines with a 7 pt gap).

  python3 transcription/forum_doc.py SPEC.json ATTRIBUTED.json --out-dir ~/Desktop/Troy/TSD

SPEC.json   {basename, title, subtitle, host_line, date_line, ballot_line, recording_line,
             questions_intro, questions: [{text, ts_ms, note?}], transcript_note}
ATTRIBUTED  {"utterances": [{start (ms), name, text}, ...]}  -- the authored-mapping output

The .docx is written with python-docx. The .pdf is printed from an HTML twin by headless
Chrome: Word's AppleScript export times out on a document this long, and Chrome's output
is what the 2022 and 2024 editions shipped.
"""
import argparse, html, json, subprocess, sys, tempfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"


def fmt(ms):
    s = int(ms) // 1000
    h, m, s = s // 3600, (s % 3600) // 60, s % 60
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def run(p, text, size=None, bold=None, italic=None, color=None):
    r = p.add_run(text)
    if size is not None: r.font.size = Pt(size)
    if bold is not None: r.bold = bold
    if italic is not None: r.italic = italic
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return r


def build_docx(spec, utts, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    sec.left_margin = sec.right_margin = Inches(1.0)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
    h1 = doc.styles["Heading 1"]; h1.font.size = Pt(14); h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string("365F91")

    for text, size, bold, italic, color in (
        (spec["title"], 20, True, None, None),
        (spec["subtitle"], 13, True, None, None),
        (spec["host_line"] + "\n" + spec["date_line"], 11, None, None, "444444"),
        (spec["ballot_line"] + "\n" + spec["recording_line"], 9, None, True, "666666"),
    ):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lines = text.split("\n")
        r = run(p, lines[0], size, bold, italic, color)
        for extra in lines[1:]:
            r.add_break(); r.add_text(extra)
    doc.add_paragraph()

    doc.add_heading("Questions Asked of the Candidates", level=1)
    run(doc.add_paragraph(), spec["questions_intro"], 10, None, True, "555555")
    for q in spec["questions"]:
        p = doc.add_paragraph(style="List Number"); p.paragraph_format.space_after = Pt(8)
        run(p, q["text"]); run(p, f"   [{fmt(q['ts_ms'])}]", 9, None, None, "888888")
        if q.get("note"):
            n = doc.add_paragraph(); n.paragraph_format.space_after = Pt(8)
            n.paragraph_format.left_indent = Pt(36)
            run(n, q["note"], 9, None, True, "888888")
    doc.add_paragraph()

    doc.add_heading("Full Transcript", level=1)
    run(doc.add_paragraph(), spec["transcript_note"], 10, None, True, "555555")
    doc.add_paragraph()
    for u in utts:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7)
        run(p, f"[{fmt(u['start'])}] ", 8.5, None, None, "999999")
        run(p, f"{u['name']}: ", None, True, None, "1F3B73")
        run(p, u["text"])
    doc.save(path)


def build_html(spec, utts):
    e = html.escape
    parts = [f"""<!doctype html><html><head><meta charset="utf-8"><title>{e(spec['basename'])}</title><style>
@page {{ size: letter; margin: 0.9in 1in; }}
body {{ font-family: Calibri, Carlito, "Segoe UI", sans-serif; font-size: 11pt; color: #000; margin: 0; }}
.c {{ text-align: center; margin: 0; }} .t1 {{ font-size: 20pt; font-weight: bold; }}
.t2 {{ font-size: 13pt; font-weight: bold; }} .t3 {{ font-size: 11pt; color: #444; }}
.t4 {{ font-size: 9pt; color: #666; font-style: italic; }}
h1 {{ font-size: 14pt; font-weight: bold; color: #365F91; margin: 18pt 0 6pt 0; }}
.intro {{ font-size: 10pt; font-style: italic; color: #555; margin: 0 0 8pt 0; }}
ol {{ margin: 0; padding-left: 22pt; }} ol li {{ margin: 0 0 8pt 0; }}
.qts {{ font-size: 9pt; color: #888; }} .note {{ font-size: 9pt; font-style: italic; color: #888; margin: 2pt 0 0 0; }}
p.u {{ margin: 0 0 7pt 0; }} p.u .ts {{ font-size: 8.5pt; color: #999; }} p.u b {{ color: #1F3B73; }}
</style></head><body>
<p class="c t1">{e(spec['title'])}</p><p class="c t2">{e(spec['subtitle'])}</p>
<p class="c t3">{e(spec['host_line'])}<br>{e(spec['date_line'])}</p>
<p class="c t4">{e(spec['ballot_line'])}<br>{e(spec['recording_line'])}</p>
<h1>Questions Asked of the Candidates</h1><p class="intro">{e(spec['questions_intro'])}</p><ol>"""]
    for q in spec["questions"]:
        note = f'<div class="note">{e(q["note"])}</div>' if q.get("note") else ""
        parts.append(f'<li>{e(q["text"])}<span class="qts">&nbsp;&nbsp;&nbsp;[{fmt(q["ts_ms"])}]</span>{note}</li>')
    parts.append(f'</ol><h1>Full Transcript</h1><p class="intro">{e(spec["transcript_note"])}</p>')
    for u in utts:
        parts.append(f'<p class="u"><span class="ts">[{fmt(u["start"])}]</span> <b>{e(u["name"])}:</b> {e(u["text"])}</p>')
    parts.append("</body></html>")
    return "".join(parts)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec"); ap.add_argument("attributed")
    ap.add_argument("--out-dir", required=True); ap.add_argument("--keep-html", action="store_true")
    a = ap.parse_args()
    spec = json.load(open(a.spec)); utts = json.load(open(a.attributed))["utterances"]
    utts = sorted(utts, key=lambda u: u["start"])
    out = Path(a.out_dir).expanduser(); out.mkdir(parents=True, exist_ok=True)
    base = out / spec["basename"]
    build_docx(spec, utts, base.with_suffix(".docx"))
    htmlp = base.with_suffix(".html") if a.keep_html else Path(tempfile.mkdtemp()) / (spec["basename"] + ".html")
    htmlp.write_text(build_html(spec, utts), encoding="utf-8")
    r = subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={base.with_suffix('.pdf')}", htmlp.as_uri()],
                       capture_output=True, text=True, timeout=180)
    if not base.with_suffix(".pdf").exists():
        sys.exit(f"chrome produced no pdf: {r.stderr[-800:]}")
    print(f"wrote {base.with_suffix('.docx').name} and .pdf in {out}  ({len(utts)} lines, {len(spec['questions'])} questions)")


if __name__ == "__main__":
    main()
